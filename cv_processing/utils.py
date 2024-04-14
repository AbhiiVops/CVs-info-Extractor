from openpyxl import Workbook
import PyPDF2
from docx import Document
from zipfile import BadZipFile
import re
import os

def extract_text_from_pdf(uploaded_file):
    reader = PyPDF2.PdfReader(uploaded_file)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_docx(uploaded_file):
    try:
        doc = Document(uploaded_file)
        text = " ".join([paragraph.text for paragraph in doc.paragraphs])
        return text
    except BadZipFile as e:
        print(f"Error processing file: {e}")
        return "" # Return an empty string or handle the error as needed

def extract_email_and_phone(text):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    phone_pattern = r'\b\d{3}[-.]?\d{3}[-.]?\d{4}\b'
    
    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)
    
    return emails, phones

def extract_additional_info(text):
    full_name_pattern = r'\b[A-Z][a-z]+ [A-Z][a-z]+\b'
    dob_pattern = r'\b(0[1-9]|[12][0-9]|3[01])[- /.](0[1-9]|1[012])[- /.](19|20)\d\d\b'
    nationality_pattern = r'\b(American|British|Canadian|Australian|Indian|Chinese|Japanese|French|German|Spanish|Italian|Russian|Brazilian|Mexican|South African|South Korean|Indonesian|Malaysian|Singaporean|Thai|Vietnamese|Philippine|Dutch|Belgian|Swiss|Greek|Turkish|Irish|Portuguese|Polish|Czech|Hungarian|Romanian|Swedish|Norwegian|Danish|Finnish|Icelandic|Slovak|Croatian|Serbian|Bulgarian|Albanian|Bosnian|Kosovar|Macedonian|Montenegrin|Armenian|Azerbaijani|Georgian|Moldovan|Ukrainian|Belarusian|Russian|Kazakh|Kyrgyz|Tajik|Turkmen|Uzbek|Afghan|Albanian|Armenian|Azerbaijani|Bangladeshi|Bhutanese|Bolivian|Bosnian|Brazilian|Bulgarian|Burmese|Cambodian|Cameroonian|Chadian|Chilean|Chinese|Colombian|Comorian|Congolese|Croatian|Cuban|Cypriot|Czech|Danish|Djiboutian|Dominican|Ecuadorian|Egyptian|Eritrean|Estonian|Ethiopian|Filipino|Fijian|Finnish|French|Gabonese|Gambian|Georgian|German|Ghanaian|Greek|Guatemalan|Guinean|Guyanese|Haitian|Honduran|Hungarian|Icelandic|Indian|Indonesian|Iranian|Iraqi|Irish|Israeli|Italian|Ivorian|Jamaican|Japanese|Jordanian|Kazakhstani|Kenyan|Kittitian|Kosovan|Kuwaiti|Kyrgyzstani|Laotian|Latvian|Lebanese|Liberian|Libyan|Lithuanian|Luxembourger|Macanese|Macedonian|Malagasy|Malawian|Malaysian|Maldivian|Malian|Maltese|Mauritanian|Mauritian|Mexican|Moldovan|Mongolian|Montenegrin|Moroccan|Mozambican|Namibian|Nepalese|Nicaraguan|Nigerian|Norwegian|Omani|Pakistani|Palauan|Panamanian|Paraguayan|Peruvian|Polish|Portuguese|Qatari|Romanian|Russian|Rwandan|Salvadoran|Saudi|Samoan|Senegalese|Serbian|Seychellois|Sierra Leonean|Singaporean|Slovak|Slovenian|Somali|South African|South Korean|Spanish|Sri Lankan|Sudanese|Surinamese|Swazi|Swedish|Swiss|Syrian|Taiwanese|Tajik|Tanzanian|Thai|Togolese|Tongan|Tunisian|Turkish|Tuvaluan|Ugandan|Ukrainian|Uruguayan|Uzbekistani|Venezuelan|Vietnamese|Yemeni|Zambian|Zimbabwean)\b'
    linkedin_pattern = r'(?:https?://)?(?:www\.)?(?:linkedin\.com/)([a-zA-Z0-9_-]+)'
    professional_summary_pattern = r'\b(Professional Summary|Career Objective|Objective|Summary):\s*([\s\S]*?)\b(Education|Skills|Experience)'
    skills_pattern = r'\b(Skills|Competencies):\s*([\s\S]*?)\b(Education|Experience)'
    education_pattern = r'\b(Education|Academic Background):\s*([\s\S]*?)\b(Skills|Experience)'

    full_name = re.search(full_name_pattern, text)
    dob = re.search(dob_pattern, text)
    nationality = re.search(nationality_pattern, text)
    linkedin = re.search(linkedin_pattern, text)
    professional_summary = re.search(professional_summary_pattern, text, re.IGNORECASE)
    skills = re.search(skills_pattern, text, re.IGNORECASE)
    education = re.search(education_pattern, text, re.IGNORECASE)

    full_name = full_name.group() if full_name else 'N/A'
    dob = dob.group() if dob else 'N/A'
    nationality = nationality.group() if nationality else 'N/A'
    linkedin = linkedin.group(1) if linkedin else 'N/A'
    professional_summary = professional_summary.group(2) if professional_summary else 'N/A'
    skills = skills.group(2) if skills else 'N/A'
    education = education.group(2) if education else 'N/A'

    return {
        'full_name': full_name,
        'dob': dob,
        'nationality': nationality,
        'linkedin': linkedin,
        'professional_summary': professional_summary,
        'skills': skills,
        'education': education
    }

def write_to_excel(data, file_path):
    wb = Workbook()
    ws = wb.active
    
    ws.append(['File Name', 'Email', 'Phone', 'Full Name', 'DOB', 'Nationality', 'LinkedIn', 'Professional Summary', 'Skills', 'Education', 'Full Text'])
    
    for item in data:
        ws.append([item['file_name'], item['email'], item['phone'], item['full_name'], item['dob'], item['nationality'], item['linkedin'], item['professional_summary'], item['skills'], item['education'], item['full_text']])
    
    wb.save(file_path)

def process_cvs(cv_files, output_file):
    data = []
    for cv_file in cv_files:
        if cv_file.name.endswith('.pdf'):
            text = extract_text_from_pdf(cv_file)
        elif cv_file.name.endswith('.doc') or cv_file.name.endswith('.docx'):
            text = extract_text_from_docx(cv_file)
        else:
            print(f"Unsupported file type: {cv_file.name}")
            continue
        
        extracted_emails, extracted_phones = extract_email_and_phone(text)
        emails = ', '.join(extracted_emails)
        phones = ', '.join(extracted_phones)
        
        additional_info = extract_additional_info(text)
        
        data.append({
            'file_name': cv_file.name,
            'email': emails,
            'phone': phones,
            'full_name': additional_info['full_name'],
            'dob': additional_info['dob'],
            'nationality': additional_info['nationality'],
            'linkedin': additional_info['linkedin'],
            'professional_summary': additional_info['professional_summary'],
            'skills': additional_info['skills'],
            'education': additional_info['education'],
            'full_text': text
        })
    
    write_to_excel(data, output_file)
